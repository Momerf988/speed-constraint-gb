"""Load the STATS19 code lists from the DfT data guide into Postgres.

Creates two tables in the raw schema:
    raw.code_list        one row per (table, column, code) with its label,
                         e.g. ('collision', 'collision_severity', '1', 'Fatal')
    raw.code_conversion  how the pre-2024 "_historic" codes map to the 2024 codes

The data guide does not always use the same column names as the CSV files, so
guide names are renamed to the CSV names here. That way every code can be
joined to the raw column it describes.

Usage:
    python db/lookups.py
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

import openpyxl
import psycopg

ROOT = Path(__file__).resolve().parent.parent
GUIDE = ROOT / "data" / "raw" / "dft-road-casualty-statistics-road-safety-open-dataset-data-guide-2025.xlsx"
DSN = os.environ.get("DATABASE_URL", "dbname=speed_gb")

DATA_TABLES = {"collision", "vehicle", "casualty"}

# Data guide column name -> CSV column name, where the two disagree.
GUIDE_TO_CSV = {
    "enhanced_collision_severity": "enhanced_severity_collision",
    "did_police_officer_attend_scene_of_collision": "did_police_officer_attend_scene_of_accident",
    "lsoa_of_collision_location": "lsoa_of_accident_location",
    "collision_adjusted_serious": "collision_adjusted_severity_serious",
    "collision_adjusted_slight": "collision_adjusted_severity_slight",
    "casualty_adjusted_serious": "casualty_adjusted_severity_serious",
    "casualty_adjusted_slight": "casualty_adjusted_severity_slight",
}

# Codes are stored as text because some are area codes such as 'E06000001'.
# Numeric raw columns are joined with a cast, e.g. l.code = c.speed_limit::text.
DDL = """
drop table if exists raw.code_list;
create table raw.code_list (
    table_name  text not null,
    field_name  text not null,
    code        text not null,
    label       text not null,
    note        text,
    primary key (table_name, field_name, code)
);

drop table if exists raw.code_conversion;
create table raw.code_conversion (
    table_name      text not null,
    historic_field  text not null,
    historic_code   text,           -- null where the 2024 code has no pre-2024 equivalent
    historic_label  text,
    current_field   text not null,
    current_code    text not null,
    current_label   text            -- the guide leaves some blank, e.g. codes mapped to -1
);
"""


def text(value) -> str | None:
    """Spreadsheet cell -> trimmed text. Empty cells and the guide's '-' become None."""
    if value is None:
        return None
    value = str(value).strip()
    return None if value in ("", "-") else value


def code_list_rows(workbook) -> tuple[list[tuple], dict[str, int]]:
    rows: dict[tuple, tuple] = {}
    stats = {"read": 0, "kept": 0, "renamed": 0, "skipped: not a data table": 0,
             "skipped: no code or no label": 0, "skipped: duplicate": 0}

    for table, field, code, label, note in workbook["2024_code_list"].iter_rows(min_row=2, values_only=True):
        stats["read"] += 1
        if table not in DATA_TABLES:  # e.g. 'historical_revisions', which documents corrections
            stats["skipped: not a data table"] += 1
            continue
        code, label = text(code), text(label)
        if code is None or label is None:  # column descriptions and formats such as '(DD/MM/YYYY)'
            stats["skipped: no code or no label"] += 1
            continue
        if field in GUIDE_TO_CSV:
            field = GUIDE_TO_CSV[field]
            stats["renamed"] += 1

        key = (table, field, code)
        if key in rows:
            # The guide repeats some rows, and lists one column under both its old and new name.
            if rows[key][3] != label:
                raise ValueError(f"{key} has two different labels: {rows[key][3]!r} and {label!r}")
            stats["skipped: duplicate"] += 1
            continue
        rows[key] = (table, field, code, label, text(note))

    stats["kept"] = len(rows)
    return list(rows.values()), stats


def conversion_rows(workbook) -> list[tuple]:
    return [
        (table, h_field, text(h_code), text(h_label), c_field, text(c_code), text(c_label))
        for table, h_field, h_code, h_label, c_field, c_code, c_label
        in workbook["2011_to_2024_conversion"].iter_rows(min_row=2, values_only=True)
    ]


def main() -> int:
    workbook = openpyxl.load_workbook(GUIDE, read_only=True)
    codes, stats = code_list_rows(workbook)
    conversions = conversion_rows(workbook)

    with psycopg.connect(DSN) as conn, conn.cursor() as cur:
        cur.execute(DDL)
        with cur.copy("copy raw.code_list (table_name, field_name, code, label, note) from stdin") as copy:
            for row in codes:
                copy.write_row(row)
        with cur.copy("copy raw.code_conversion from stdin") as copy:
            for row in conversions:
                copy.write_row(row)

    print("raw.code_list")
    for name, count in stats.items():
        print(f"    {name:32s} {count:,}")
    print(f"raw.code_conversion\n    rows                             {len(conversions):,}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
